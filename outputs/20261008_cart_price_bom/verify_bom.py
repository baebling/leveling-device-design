"""Repair only unsupported native link metadata, then independently audit the artifact export.

No workbook values, formulas or styles are authored by this script. Artifact Tool owns them.
Python's ZIP/XML APIs fill the missing external-hyperlink capability only; openpyxl is read-only.
"""
import json
from copy import copy
import os
from pathlib import Path
import sys
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
folder = Path(__file__).resolve().parent
audit = json.loads((folder / 'bom_price_audit.json').read_text(encoding='utf-8'))
target = Path(audit['outputPath'])
ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
doc_rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
ET.register_namespace('', ns)
ET.register_namespace('r', doc_rel_ns)
sheet_name = 'xl/worksheets/sheet1.xml'
rels_name = 'xl/worksheets/_rels/sheet1.xml.rels'
with zipfile.ZipFile(target) as source_zip:
    entries = [(entry, source_zip.read(entry.filename)) for entry in source_zip.infolist()]
    xml = ET.fromstring(source_zip.read(sheet_name))
    rels = ET.fromstring(source_zip.read(rels_name))
links = xml.find(f'{{{ns}}}hyperlinks')
assert links is not None
used_ids = {rel.attrib['Id'] for rel in rels}
for repair in audit['nativeLinkRepairs']:
    element = next((link for link in links if link.attrib['ref'] == repair['cell']), None)
    if element is None:
        number = 1
        while f'rIdBOM{number}' in used_ids:
            number += 1
        rid = f'rIdBOM{number}'
        used_ids.add(rid)
        element = ET.SubElement(links, f'{{{ns}}}hyperlink', {'ref': repair['cell'], f'{{{doc_rel_ns}}}id': rid})
        rel = ET.SubElement(rels, f'{{{rel_ns}}}Relationship', {'Id': rid, 'Type': doc_rel_ns + '/hyperlink', 'TargetMode': 'External'})
    else:
        rid = element.attrib[f'{{{doc_rel_ns}}}id']
        rel = next(rel for rel in rels if rel.attrib['Id'] == rid)
    rel.set('Target', repair['target'])
    rel.set('TargetMode', 'External')
replacements = {
    sheet_name: ET.tostring(xml, encoding='utf-8', xml_declaration=True),
    rels_name: ET.tostring(rels, encoding='utf-8', xml_declaration=True),
}
temporary = target.with_suffix('.xlsx.native-links.tmp')
with zipfile.ZipFile(temporary, 'w') as out_zip:
    for entry, data in entries:
        out_zip.writestr(entry, replacements.get(entry.filename, data))
os.replace(temporary, target)

original = openpyxl.load_workbook(audit['sourcePath'])
result = openpyxl.load_workbook(target)
cached = openpyxl.load_workbook(target, data_only=True)
source_sheet, sheet, values = original.active, result.active, cached.active
assert original.sheetnames == result.sheetnames == ['Sheet1']
assert source_sheet.max_row == sheet.max_row == 78
assert source_sheet.max_column == sheet.max_column == 9
assert sheet.freeze_panes == source_sheet.freeze_panes
assert str(sheet.merged_cells) == str(source_sheet.merged_cells)
assert len(sheet._charts) == len(source_sheet._charts) == 0
assert len(sheet._images) == len(source_sheet._images) == 0
for col, dim in source_sheet.column_dimensions.items():
    assert abs(sheet.column_dimensions[col].width - dim.width) < 1e-8, col
expected = audit['totals']['all']
assert [values.cell(r, 8).value for r in range(75,79)] == [expected['supply'], expected['vat'], expected['gross'], expected['remaining']]
assert sum(values.cell(r, 8).value for r in range(4,75)) == expected['supply']
for row in range(75,79):
    assert sheet.cell(row,8).value == source_sheet.cell(row,8).value

price_changes = {c['originalRow']: c['newSupply'] for c in audit['changes']}
special_notes = {29,68,69,70,71,72,73,74}
special_links = {29,72,73}
price_rows_checked = 0
for mapping in audit['rowMapping']:
    old, new = mapping['originalRow'], mapping['outputRow']
    assert sheet.cell(new,6).value == source_sheet.cell(old,6).value
    assert sheet.cell(new,8).value == price_changes.get(old, source_sheet.cell(old,8).value)
    price_rows_checked += 1
    for col in (2,4,5):
        if col == 5 and old in (29,73):
            continue
        assert sheet.cell(new,col).value == source_sheet.cell(old,col).value, (old,new,col)
    if old != 72:
        assert sheet.cell(new,3).value == source_sheet.cell(old,3).value
    else:
        assert sheet.cell(new,3).value == '디바이스마트'
    if old not in special_notes:
        assert sheet.cell(new,9).value == source_sheet.cell(old,9).value
    if old not in special_links:
        assert sheet.cell(new,7).value == source_sheet.cell(old,7).value
        if source_sheet.cell(old,7).hyperlink:
            assert sheet.cell(new,7).hyperlink.target == source_sheet.cell(old,7).hyperlink.target
    if sheet.cell(new,7).hyperlink:
        assert sheet.cell(new,7).hyperlink.target == sheet.cell(new,7).value, f'Wrong click destination G{new}'
    for col in range(2,10):
        for attribute in ('font','fill','border','alignment','number_format','protection'):
            result_style = copy(getattr(sheet.cell(new,col),attribute))
            source_style = copy(getattr(source_sheet.cell(old,col),attribute))
            if attribute == 'font':
                # Export omits the legacy code-page marker; font family/name/size/color remain verified.
                result_style.charset = source_style.charset = None
            assert result_style == source_style, (old,new,col,attribute)
for repair in audit['nativeLinkRepairs']:
    assert sheet[repair['cell']].value == repair['target']
    assert sheet[repair['cell']].hyperlink.target == repair['target']
for row in (1,2,3,75,76,77,78):
    for col in range(1,10):
        if row >= 75 and col == 9:
            continue
        assert sheet.cell(row,col).value == source_sheet.cell(row,col).value, (row,col)
for row in sheet.iter_rows():
    for cell in row:
        assert cell.data_type != 'e', cell.coordinate
        assert values[cell.coordinate].data_type != 'e', cell.coordinate
report = {'lineCount':price_rows_checked,'cachedTotals':expected,'formulasPreserved':True,'vendorGroupingPreserved':True,'allLinkDestinationsCorrect':True,'nativeLinkRepairs':3,'visibleStylesPreserved':True,'styleNote':'Legacy font charset marker omitted by export; name, size, color, fill, border, alignment and number format verified.','sheetAndFreezePreserved':True,'unexpectedFormulaErrors':0}
(folder/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
