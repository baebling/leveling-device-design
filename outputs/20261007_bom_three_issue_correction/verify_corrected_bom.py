"""Independently read exported XLSX and check purchase/assembly corrections.

Spreadsheet content/layout/formulas are authored with artifact-tool. The only
ZIP post-processing attaches native hyperlinks to the already visible URLs;
openpyxl is used strictly for read-only validation, never workbook authoring.
"""
from collections import Counter
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
XL = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'


def attach_exact_hyperlinks(target, rows):
    with ZipFile(target) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    ET.register_namespace('', XL)
    ET.register_namespace('r', REL)
    sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    old = sheet.find(f'{{{XL}}}hyperlinks')
    if old is not None:
        sheet.remove(old)
    links = ET.Element(f'{{{XL}}}hyperlinks')
    relationship_path = 'xl/worksheets/_rels/sheet1.xml.rels'
    rels = ET.fromstring(parts.get(relationship_path, f'<Relationships xmlns="{PKG}"/>'.encode()))
    for element in list(rels):
        if element.get('Type', '').endswith('/hyperlink'):
            rels.remove(element)
    for i, row in enumerate(rows, 4):
        url = row[5]
        assert isinstance(url, str) and url.startswith('https://')
        identifier = f'rIdBomLink{i}'
        ET.SubElement(links, f'{{{XL}}}hyperlink', {'ref':f'G{i}', f'{{{REL}}}id':identifier})
        ET.SubElement(rels, f'{{{PKG}}}Relationship', {'Id':identifier, 'Type':REL+'/hyperlink', 'Target':url, 'TargetMode':'External'})
    after_links = {'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','drawing','legacyDrawing','extLst'}
    index = next((i for i, child in enumerate(sheet) if child.tag.split('}')[-1] in after_links), len(sheet))
    sheet.insert(index, links)
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(sheet, encoding='utf-8', xml_declaration=True)
    parts[relationship_path] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
    temporary = target.with_suffix('.linkrepair.xlsx')
    with ZipFile(temporary, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    temporary.replace(target)


def validate():
    manifest = json.loads((HERE / 'bom_change_manifest.json').read_text(encoding='utf-8'))
    target = Path(manifest['outputPath'])
    attach_exact_hyperlinks(target, manifest['rows'])
    original = load_workbook(manifest['sourcePath'], data_only=True).active
    workbook = load_workbook(target, data_only=True)
    sheet = workbook.active
    formula_sheet = load_workbook(target, data_only=False).active
    end = manifest['subtotalRow'] - 1
    actual = [[sheet.cell(r, c).value for c in range(2,10)] for r in range(4, end+1)]
    assert actual == manifest['rows'], 'Exported rows differ from authored purchase data'
    originals = [[original.cell(r,c).value for c in range(2,10)] for r in range(4,67)]
    changed = {c['oldModel']:c for c in manifest['changes']}
    indexed = {r[3]:r for r in actual}
    unexpected = []
    for before in originals:
        if before[3] in changed:
            record = changed[before[3]]
            assert before == record['before'], f'Original changed row mismatch {before[3]}'
            assert indexed[record['after'][3]] == record['after']
        elif indexed.get(before[3]) != before:
            unexpected.append(before[3])
    assert not unexpected
    links_match = all(sheet[f'G{r}'].hyperlink and sheet[f'G{r}'].hyperlink.target == sheet[f'G{r}'].value for r in range(4,end+1))
    assert links_match
    assert not any('/result/' in r[5] or '/s/?q=' in r[5] or '/s-1/' in r[5] for r in actual), 'Search-only purchasing link remains'
    assert len(indexed) == len(actual), 'Duplicate purchase model'
    assert all(isinstance(r[4],int) and r[4]>0 and isinstance(r[6],(int,float)) and r[6]>0 for r in actual)
    blocks = [r[1] for i,r in enumerate(actual) if i==0 or r[1]!=actual[i-1][1]]
    assert len(blocks)==len(set(blocks)), 'Vendor is split across blocks'
    total = sum(r[6] for r in actual)
    assert total == sheet[f'H{manifest["subtotalRow"]}'].value == manifest['supply']
    assert formula_sheet[f'H{manifest["subtotalRow"]}'].value == f'=SUM(H4:H{end})'
    assert not any(cell.data_type == 'e' for row in sheet for cell in row)
    assert list(str(rng) for rng in sheet.merged_cells.ranges)==['B1:I1'], 'Unexpected covering merge'
    for row in range(4,end+1):
        assert sheet.row_dimensions[row].height >= 72
    pin = json.loads((HERE/'pin_stack_audit.json').read_text(encoding='utf-8'))
    sensor = json.loads((HERE/'sensor_mount_audit.json').read_text(encoding='utf-8'))
    assert pin['pose_count']==sensor['pose_count']==27
    assert pin['invalid_boolean_count']==sensor['invalid_boolean_count']==0
    assert pin['unexpected_interference_count']==sensor['unexpected_interference_count']==0
    assert indexed['HCDGH6-40'][4] == 3 and 'HCDGH6-35' not in indexed
    assert indexed['WSSB10-6-4'][4] == 6
    assert indexed['CIMR6-10-0.1'][4] == 9 and indexed['CIMR6-10-0.5'][4] == 12
    assert indexed['ENBT-100-100-5'][4] == sensor['plate_quantity'] == 2
    assert indexed['CB6-25'][4] == sensor['m6_profile_fasteners'] == 2
    assert indexed['CB8-25'][4] == sensor['m8_profile_fasteners'] == 2
    assert indexed['WSSB10-6-5'][4]+indexed['WSSB12-8-5'][4] == sensor['spacer_5mm_quantity'] == 8
    assert indexed['E-SPN306'][4] == 16+6+sensor['m6_profile_fasteners']
    # Existing M4/M8 nut packages are retained; do not invent quantities.
    assert indexed['K14364804'][4] == indexed['K06130054'][4] == indexed['K14671419'][4] == 1
    evidence = {
        'main_splice': {'model':indexed['WAGO 221-615 / 14954324'][3], 'minimum_mm2':.5,'maximum_mm2':6.,
            'rated_current_a':30., 'requires_crimp_tool':False, 'used_quantity':2,'purchase_quantity':5,
            'wire_cross_sections_mm2':[2.08,4.], 'protected_path_a':20,
            'technical_source':'https://www.wago.com/global/electrical-interconnections/discover-installation-terminal-blocks-and-connectors/221'},
        'control_splice': {'model':indexed['WAGO 221-415 / 14954315 / 25EA'][3], 'minimum_mm2':.14,'maximum_mm2':4.,'control_wire_mm2':.205},
        'pin_stack':pin['stack'],
        'sensor_mount':{k:v for k,v in sensor.items() if k!='measurements'},
        'bom': {'unrelated_row_changes':unexpected,'purchase_rows':len(actual), 'native_hyperlinks':len(actual),
            'independent_supply_sum':total,'spreadsheet_supply_sum':sheet[f'H{manifest["subtotalRow"]}'].value,
            'vat_estimate':manifest['vat'],'vat_included_estimate':manifest['total'],'budget_remaining_estimate':manifest['remaining'],
            'supply_delta_from_previous':total-sum(r[6] for r in originals),
            'links_match_visible_urls':links_match,'vendor_blocks_contiguous':True,
            'formula_errors':0,'unexpected_merges':0,'delivery_cost_included':False},
        'scope': {'new_metal_machining':False,'existing_meviy_metal_shapes_changed':False,'actual_powered_test_performed':False,
            'moving_cable_sweep_verified':False,'750N_strength_release':False},
    }
    (HERE/'verification.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(evidence,ensure_ascii=False,indent=2))


if __name__=='__main__':
    validate()
