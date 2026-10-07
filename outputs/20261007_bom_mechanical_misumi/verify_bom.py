"""Independent read-only workbook checks plus native-hyperlink XML repair."""
import hashlib
import importlib.util
import json
from pathlib import Path
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
m=json.loads((HERE/'bom_change_manifest.json').read_text(encoding='utf-8'))
helper=HERE.parent/'20261007_bom_three_issue_correction'/'verify_corrected_bom.py'
spec=importlib.util.spec_from_file_location('hyperlink_helper',helper)
module=importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
target=Path(m['outputPath'])
module.attach_exact_hyperlinks(target,m['rows'])
s=load_workbook(target,data_only=True).active
f=load_workbook(target,data_only=False).active
actual=[[s.cell(r,c).value for c in range(2,10)] for r in range(4,75)]
assert actual==m['rows']
changes={c['oldModel']:c for c in m['changes']}
indexed={r[3]:r for r in actual}
assert len(indexed)==len(actual)==71
for before in m['originals']:
    after=changes.get(before[3],{}).get('after',before)
    assert indexed[after[3]]==after,before[3]
assert len(changes)==11
assert sum(c['vendorChanged'] for c in m['changes'])==7
assert sum(r[1]=='나비엠알오' for r in actual)==7
assert sum(r[1]=='나비엠알오' and r[7].startswith(('S','F')) for r in actual)==4
blocks=[r[1] for i,r in enumerate(actual) if i==0 or r[1]!=actual[i-1][1]]
assert len(blocks)==len(set(blocks))
assert all(s[f'G{r}'].hyperlink.target==s[f'G{r}'].value for r in range(4,75))
assert all('/result/' not in r[5] for r in actual)
assert sum(r[6] for r in actual)==s['H75'].value==m['supply']
assert s['H76'].value==int(m['supply']*.1+.5)==m['vat']
assert s['H77'].value==m['total']==m['supply']+m['vat']
assert s['H78'].value==4000000-m['total']
assert [f[f'H{r}'].value for r in range(75,79)]==['=SUM(H4:H74)','=ROUND(H75*10%,0)','=SUM(H75:H76)','=$I$2-H77']
assert not any(cell.data_type=='e' for row in s for cell in row)
assert [str(r) for r in s.merged_cells.ranges]==['B1:I1']
assert all(s.row_dimensions[r].height>=92 for r in range(4,75))
assert all(isinstance(r[4],int) and r[4]>0 and r[6]>0 for r in actual)
for c in m['changes']:
    if c['vendorChanged']:
        e=c['evidence']
        assert c['after'][4]==e['quantity']>=e['minimumOrder']
        assert c['after'][6]==e['unitSupply']*e['quantity']
        if 'previousPhysicalQuantity' in e:
            assert e['quantity']>=e['previousPhysicalQuantity']
assert indexed['E-DNF4040BK-700'][4]==2
assert indexed['E-DNF4040BK-620'][4]==4
assert indexed['E-DCBK4035'][4]==8
assert indexed['E-SPN408'][4]==indexed['CB8-16'][4]==indexed['CB8-20'][4]==100
assert indexed['HXN-ST-M8-12'][4]==6
assert indexed['K06130485'][4]==indexed['K14364804'][4]==indexed['K06130054'][4]==1
assert indexed['HCDGH6-40'][4]==3 and indexed['E-SPN306'][4]==24
assert indexed['CB8-25'][4]==indexed['CB6-25'][4]==2
assert hashlib.sha256(Path(m['sourcePath']).read_bytes()).hexdigest()==m['sourceHash']
old_total=m['previousSupply']+int(m['previousSupply']*.1+.5)
result={'purchase_rows':71,'reviewed_mechanical_rows':11,'misumi_vendor_changes':7,'navimro_remaining_mechanical_rows':4,'navimro_remaining_electrical_rows':3,'native_hyperlinks':71,'unchanged_rows_preserved':60,'formula_errors':0,'vendor_blocks_contiguous':True,'source_preserved':True,'supply':m['supply'],'vat':m['vat'],'total':m['total'],'vat_included_reduction':old_total-m['total'],'remaining_budget':m['remaining'],'vendor_totals':m['vendorTotals'],'cart_changed':False,'new_metal_machining':False,'actual_powered_test':False,'live_prices_not_final_checkout':True}
(HERE/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=True))
