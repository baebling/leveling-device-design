"""Check the buyer-visible RevE workbook, including actual Excel link targets.

Usage: python tests/check_reve_bom_xlsx.py path/to/workbook.xlsx
"""

from pathlib import Path
import sys
from zipfile import ZipFile

from openpyxl import load_workbook


INTENDED_VALUE_EDITS = {
    "B1", "D18", "I17", "I18", "I24", "D38", "I38", "D43", "I43",
    "D44", "I44", "I47", "I52", "I53", "D56", "E56", "I56", "I58",
    "I59", "I60", "H66", "I66", "B71", "B73", "B74", "I74",
}


def check(path: Path, source: Path | None = None) -> None:
    with ZipFile(path) as archive:
        assert archive.testzip() is None, "Corrupt XLSX ZIP member"
    values_book = load_workbook(path, data_only=True)
    formula_book = load_workbook(path, data_only=False)
    sheet = formula_book["Sheet1"]
    values_sheet = values_book["Sheet1"]

    expected_rows = range(4, 71)
    ids = []
    link_errors = []
    numeric_total = 0
    supplier_groups = []
    for row in expected_rows:
        supplier = sheet[f"C{row}"].value
        assert isinstance(supplier, str) and supplier, f"Missing seller C{row}"
        if not supplier_groups or supplier_groups[-1] != supplier:
            supplier_groups.append(supplier)
        assert sheet[f"D{row}"].value and sheet[f"E{row}"].value, (
            f"Missing part name/model on row {row}"
        )
        quantity = sheet[f"F{row}"].value
        assert isinstance(quantity, (int, float)) and quantity > 0, (
            f"Invalid quantity F{row}: {quantity}"
        )
        visible = sheet[f"G{row}"].value
        actual = (
            sheet[f"G{row}"].hyperlink.target
            if sheet[f"G{row}"].hyperlink is not None
            else None
        )
        if not visible or visible != actual:
            link_errors.append((row, visible, actual))
        note = str(sheet[f"I{row}"].value or "")
        ids.append(note.split(".", 1)[0])
        cost = sheet[f"H{row}"].value
        assert isinstance(cost, (int, float)) and cost > 0, (
            f"Missing or nonpositive amount H{row}: {cost}"
        )
        numeric_total += cost

    assert len(ids) == 67, f"Expected 67 BOM lines, found {len(ids)}"
    assert len(set(ids)) == 67, "Duplicate or missing BOM IDs"
    assert len(supplier_groups) == len(set(supplier_groups)), "Seller grouping is broken"
    assert not link_errors, f"{len(link_errors)} bad hyperlink rows: {link_errors[:5]}"
    assert sheet["H71"].value == "=SUM(H4:H70)", "Subtotal formula moved"
    assert values_sheet["H71"].value == numeric_total, (
        f"Cached subtotal {values_sheet['H71'].value} != source sum {numeric_total}"
    )
    assert values_sheet["H72"].value == (numeric_total + 5) // 10, "VAT cache error"
    assert values_sheet["H73"].value == numeric_total + values_sheet["H72"].value, (
        "Gross total cache error"
    )
    assert values_sheet["H74"].value == 4_000_000 - values_sheet["H73"].value, (
        "Budget balance cache error"
    )
    assert sheet["H66"].value == 600, "Three M6x25 studs must have priced 600 KRW supply"
    assert "발주 보류" in str(sheet["B1"].value), "Purchase-hold banner missing"
    assert "LM4075OE-1075" in str(sheet["E56"].value), "Actuator model mismatch"
    assert "1000035578" in str(sheet["G56"].value), "Actuator SKU mismatch"
    for row in sheet:
        for cell in row:
            assert cell.data_type != "e", f"Excel error cell {cell.coordinate}: {cell.value}"
    if source is not None:
        original_book = load_workbook(source, data_only=False)
        original_sheet = original_book["Sheet1"]
        unexpected = []
        for row in range(1, max(sheet.max_row, original_sheet.max_row) + 1):
            for col in range(1, max(sheet.max_column, original_sheet.max_column) + 1):
                new = sheet.cell(row, col)
                old = original_sheet.cell(row, col)
                if new.value != old.value and new.coordinate not in INTENDED_VALUE_EDITS:
                    unexpected.append((new.coordinate, old.value, new.value))
        assert not unexpected, f"Unexpected changed cells: {unexpected[:10]}"
        original_book.close()
    values_book.close()
    formula_book.close()
    print(f"PASS: 67 IDs, 67 exact links, subtotal {numeric_total:,} KRW")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Pass output XLSX path and optional source XLSX path")
    check(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) == 3 else None)
