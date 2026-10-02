"""Behavioral tests for the Rev F buyer-review workbook checker."""

from pathlib import Path
from copy import deepcopy
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

from tests import check_revf_bom_xlsx as checker
from tests.check_revf_bom_xlsx import read_rows


ROOT = Path(__file__).resolve().parents[1]
CHECKER = ROOT / "tests" / "check_revf_bom_xlsx.py"
REV_E = (
    ROOT
    / "outputs"
    / "20261002_reve_followthrough"
    / "2026_BIZ-Lab_재료비관리_RevE_연속검수.xlsx"
)
REV_F = (
    ROOT
    / "outputs"
    / "20261002_reve_upper_pocket_bom"
    / "2026_BIZ-Lab_재료비관리_RevF_상부포켓검토.xlsx"
)
REMOVED_UPPER_IDS = {"F08A", "F08B", "F11", "F12"}


def review_rows():
    rows = [deepcopy(row) for row in read_rows(REV_E) if row["id"] not in REMOVED_UPPER_IDS]
    insertion = next(index for index, row in enumerate(rows) if row["id"] == "M03") + 1
    brackets = [
        {
            "id": f"UP01{axis}",
            "seller": "한국미스미",
            "quantity_pieces": 1,
            "link_text": None,
            "hyperlink_target": None,
            "price_krw": "견적 미확정",
            "row": 100 + index,
        }
        for index, axis in enumerate("ABC")
    ]
    return rows[:insertion] + brackets + rows[insertion:]


class RevFBomCheckerTests(unittest.TestCase):
    def test_exported_review_workbook_passes(self):
        checker.check(REV_F)
        book = load_workbook(REV_F)
        self.assertIsNone(book["Sheet1"]["G70"].value)
        self.assertIsNone(book["Sheet1"]["G70"].hyperlink)
        book.close()

    def test_summary_formula_must_not_be_stale(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tampered.xlsx"
            book = load_workbook(REV_F)
            book["Sheet1"]["H70"] = "=SUM(H4:H70)"
            book.save(path)
            book.close()
            with self.assertRaisesRegex(AssertionError, "summary formula"):
                checker.check(path)

    def test_retained_product_cannot_disappear(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "tampered.xlsx"
            book = load_workbook(REV_F)
            book["Sheet1"]["I62"] = "No ID"
            book.save(path)
            book.close()
            with self.assertRaisesRegex(AssertionError, "retained ID"):
                checker.check(path)

    def test_review_rows_pass_row_level_validation(self):
        self.assertTrue(hasattr(checker, "validate_rows"))
        checker.validate_rows(review_rows())

    def test_duplicate_id_is_rejected(self):
        rows = review_rows()
        rows.insert(1, deepcopy(rows[0]))
        with self.assertRaisesRegex(AssertionError, "Duplicate"):
            checker.validate_rows(rows)

    def test_broken_exact_product_link_is_rejected(self):
        rows = review_rows()
        rows[0]["hyperlink_target"] = "https://invalid.example/other"
        with self.assertRaisesRegex(AssertionError, "hyperlink"):
            checker.validate_rows(rows)

    def test_split_seller_group_is_rejected(self):
        rows = review_rows()
        rows.append(rows.pop(0))
        with self.assertRaisesRegex(AssertionError, "seller group"):
            checker.validate_rows(rows)

    def test_zero_bracket_quote_is_rejected(self):
        rows = review_rows()
        next(row for row in rows if row["id"] == "UP01A")["price_krw"] = 0
        with self.assertRaisesRegex(AssertionError, "quote"):
            checker.validate_rows(rows)

    def test_f10_cannot_be_removed_while_f21_is_present(self):
        rows = [row for row in review_rows() if row["id"] != "F10"]
        with self.assertRaisesRegex(AssertionError, "F10"):
            checker.validate_rows(rows)

    def test_old_upper_pivot_row_cannot_remain_active(self):
        rows = review_rows()
        rows.insert(2, deepcopy(next(row for row in read_rows(REV_E) if row["id"] == "F11")))
        with self.assertRaisesRegex(AssertionError, "obsolete"):
            checker.validate_rows(rows)

    def test_zero_numeric_candidate_price_is_rejected(self):
        rows = review_rows()
        rows[0]["price_krw"] = 0
        with self.assertRaisesRegex(AssertionError, "price"):
            checker.validate_rows(rows)

    def test_reader_converts_m03_two_packs_to_four_pieces(self):
        rows = read_rows(REV_E)
        self.assertEqual(67, len(rows))
        m03 = next(row for row in rows if row["id"] == "M03")
        self.assertEqual(4, m03["quantity_pieces"])

    def test_reader_converts_critical_fastener_packs_to_pieces(self):
        rows = {row["id"]: row for row in read_rows(REV_E)}
        self.assertEqual(57, rows["F02"]["quantity_pieces"])
        self.assertEqual(100, rows["F07"]["quantity_pieces"])
        self.assertEqual(100, rows["F10"]["quantity_pieces"])

    def test_rev_e_cannot_pass_without_three_upper_pocket_brackets(self):
        result = subprocess.run(
            [sys.executable, str(CHECKER), str(REV_E)],
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertNotEqual(0, result.returncode)
        self.assertIn("UP01A", result.stderr + result.stdout)


if __name__ == "__main__":
    unittest.main()
