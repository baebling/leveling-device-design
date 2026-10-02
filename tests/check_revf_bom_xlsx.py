"""Read-only checks for the Rev F upper-pocket purchase-review workbook."""

from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import re
import sys

from openpyxl import load_workbook


ID_PREFIX = re.compile(r"^([A-Z][A-Z0-9]*)\.")
BRACKET_IDS = {"UP01A", "UP01B", "UP01C"}
PACK_SIZES = {"F02": 57, "F07": 100, "F10": 100, "M03": 2}
OBSOLETE_UPPER_IDS = {"F08A", "F08B", "F11", "F12"}
ROOT = Path(__file__).resolve().parents[1]
REV_E = ROOT / "outputs/20261002_reve_followthrough/2026_BIZ-Lab_재료비관리_RevE_연속검수.xlsx"


def read_rows(path: Path) -> list[dict[str, object]]:
    book = load_workbook(path, data_only=False, read_only=False)
    sheet = book["Sheet1"]
    rows = []
    for row in range(4, sheet.max_row + 1):
        note = str(sheet[f"I{row}"].value or "")
        match = ID_PREFIX.match(note)
        if not match:
            continue
        link = sheet[f"G{row}"]
        rows.append(
            {
                "id": match.group(1),
                "seller": sheet[f"C{row}"].value,
                "quantity_pieces": sheet[f"F{row}"].value * PACK_SIZES.get(match.group(1), 1),
                "link_text": link.value,
                "hyperlink_target": link.hyperlink.target if link.hyperlink else None,
                "price_krw": sheet[f"H{row}"].value,
                "row": row,
            }
        )
    book.close()
    return rows


def validate_rows(rows: list[dict[str, object]]) -> None:
    """Reject silent scope, link, grouping, and pack-count regressions."""
    identifiers = [str(row["id"]) for row in rows]
    assert len(identifiers) == len(set(identifiers)), "Duplicate BOM ID"
    ids = set(identifiers)
    missing = BRACKET_IDS - ids
    assert not missing, f"Missing upper pocket bracket IDs: {', '.join(sorted(missing))}"
    obsolete = OBSOLETE_UPPER_IDS & ids
    assert not obsolete, f"Active obsolete upper-pivot rows: {', '.join(sorted(obsolete))}"
    assert "F10" in ids or "F21" not in ids, "F10 is required by F21 / G9EA"

    for row in rows:
        identifier = str(row["id"])
        price = row["price_krw"]
        if identifier in BRACKET_IDS:
            assert row["seller"] == "한국미스미", f"{identifier}: expected quote-candidate group"
            assert row["quantity_pieces"] == 1, f"{identifier}: exactly one bracket required"
            assert price == "견적 미확정", f"{identifier}: quote is not established"
            assert not row["hyperlink_target"], f"{identifier}: no exact quote/product hyperlink yet"
            continue
        assert isinstance(price, (int, float)) and price > 0, f"{identifier}: invalid price"
        assert isinstance(row["link_text"], str) and row["link_text"].startswith("https://"), (
            f"{identifier}: missing exact product link text"
        )
        assert row["hyperlink_target"] == row["link_text"], (
            f"{identifier}: hyperlink target does not match displayed exact product URL"
        )

    seen_sellers: set[str] = set()
    previous_seller = None
    for row in rows:
        seller = str(row["seller"])
        if seller != previous_seller:
            assert seller not in seen_sellers, f"Split seller group: {seller}"
            seen_sellers.add(seller)
            previous_seller = seller

    quantities = {str(row["id"]): row["quantity_pieces"] for row in rows}
    for identifier, expected in {"F02": 57, "F07": 100, "F10": 100, "M03": 4}.items():
        assert quantities.get(identifier) == expected, (
            f"{identifier}: pack/piece conversion must be {expected} pieces"
        )


def check(path: Path) -> None:
    rows = read_rows(path)
    validate_rows(rows)
    original_rows = read_rows(REV_E)
    retained = [row for row in original_rows if row["id"] not in OBSOLETE_UPPER_IDS]
    expected_ids = [str(row["id"]) for row in retained]
    insertion = expected_ids.index("M03") + 1
    expected_ids[insertion:insertion] = ["UP01A", "UP01B", "UP01C"]
    assert [row["id"] for row in rows] == expected_ids, "Missing or reordered retained ID"

    original_by_id = {row["id"]: row for row in retained}
    for row in rows:
        identifier = row["id"]
        if identifier in BRACKET_IDS:
            continue
        source = original_by_id[identifier]
        for field in ("seller", "quantity_pieces", "link_text", "hyperlink_target", "price_krw"):
            assert row[field] == source[field], f"{identifier}: retained {field} changed"

    formula_book = load_workbook(path, data_only=False)
    value_book = load_workbook(path, data_only=True)
    sheet = formula_book["Sheet1"]
    values = value_book["Sheet1"]
    last_item_row = max(int(row["row"]) for row in rows)
    assert last_item_row == 69, "Unexpected BOM row boundary"
    expected_formulas = {
        "H70": "=SUM(H4:H69)",
        "H71": "=ROUND(H70*10%,0)",
        "H72": "=SUM(H70:H71)",
        "H73": "=$I$2-H72",
    }
    for cell, expected in expected_formulas.items():
        assert sheet[cell].value == expected, f"{cell}: summary formula changed"
    assert sheet["I2"].value == 4_000_000, "Budget cap changed"
    supply = sum(int(row["price_krw"]) for row in rows if row["id"] not in BRACKET_IDS)
    vat = int((Decimal(supply) * Decimal("0.1")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    expected_values = {"H70": supply, "H71": vat, "H72": supply + vat, "H73": 4_000_000 - supply - vat}
    for cell, expected in expected_values.items():
        assert values[cell].value == expected, f"{cell}: cached summary value differs from independent calculation"
    assert "최종 총액 아님" in str(sheet["B72"].value), "Partial total must not be labeled final"
    assert "구매 가능액 아님" in str(sheet["B73"].value), "Remaining cap must not imply purchase release"
    assert sheet["G70"].value is None and sheet["G70"].hyperlink is None, "Stray summary hyperlink"
    assert "필수 미선정 항목" in str(sheet["B75"].value), "Missing unselected hardware section"
    formula_book.close()
    value_book.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Pass Rev F XLSX path")
    check(Path(sys.argv[1]))
