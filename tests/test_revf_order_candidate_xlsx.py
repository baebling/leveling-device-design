from pathlib import Path
import unittest

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK = (
    ROOT
    / "outputs"
    / "20261003_revf_order_candidate"
    / "2026_BIZ-Lab_재료비관리_RevF_발주후보.xlsx"
)


class RevFOrderCandidateWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # artifact-tool writes valid cell records without relying on the optional
        # worksheet dimension hint, so use normal mode and let openpyxl derive it.
        cls.wb = load_workbook(BOOK, data_only=False, read_only=False)
        cls.ws = cls.wb["Sheet1"]
        cls.rows = [
            [cls.ws.cell(row, col).value for col in range(2, 10)]
            for row in range(4, cls.ws.max_row + 1)
            if cls.ws.cell(row, 2).value == "재료비 구매"
        ]

    def test_required_finalized_models_exist(self):
        text = "\n".join(" | ".join(str(value or "") for value in row) for row in self.rows)
        for token in (
            "HCDGH6-35",
            "WSSB10-6-4",
            "WSSB10-6-1.5",
            "CBS6-12",
            "E-DNF3030",
            "E-SPN306",
            "E-DCBK3025",
            "22.32.0.024.4320",
            "01550300Z",
            "DKD35-P001",
            "DDR-15G-12",
            "BC-AGP-507025",
            "5070P",
        ):
            self.assertIn(token, text)

    def test_obsolete_or_incompatible_items_removed(self):
        text = "\n".join(" | ".join(str(value or "") for value in row) for row in self.rows)
        for token in (
            "G9EA-1-B",
            "MY2N-D2-GS",
            "BPS100A-M6/4*M5",
            "UL1283 8AWG",
            "JTRN-4-1M",
            "K07636024",
            "K56842696",
            "K42296106",
            "K42296071",
        ):
            self.assertNotIn(token, text)

    def test_no_unselected_or_hold_purchase_rows(self):
        for row in self.rows:
            joined = " | ".join(str(value or "") for value in row)
            self.assertNotIn("미선정", joined)
            self.assertNotIn("주문 금지", joined)
            self.assertNotIn("HOLD", joined)

    def test_custom_brackets_are_the_only_unpriced_purchase_rows(self):
        unpriced = [row for row in self.rows if not isinstance(row[6], (int, float))]
        self.assertEqual(len(unpriced), 3)
        self.assertTrue(all("상부 포켓 브래킷" in str(row[2]) for row in unpriced))
        self.assertTrue(all("사용자 견적" in str(row[6]) for row in unpriced))

    def test_each_vendor_is_one_contiguous_block(self):
        vendors = [str(row[1]) for row in self.rows]
        for vendor in set(vendors):
            indices = [index for index, value in enumerate(vendors) if value == vendor]
            self.assertEqual(indices, list(range(min(indices), max(indices) + 1)))

    def test_summary_formulas_exist(self):
        labels = {
            self.ws.cell(row, 2).value: self.ws.cell(row, 8).value
            for row in range(4, self.ws.max_row + 1)
        }
        self.assertTrue(str(labels["공급가 부분합 (상부 브래킷 견적 제외)"]).startswith("=SUM("))
        self.assertTrue(str(labels["부가세 10% (상부 브래킷 견적 제외)"]).startswith("=ROUND("))
        self.assertTrue(str(labels["VAT 포함 합계 (상부 브래킷 견적 제외)"]).startswith("=SUM("))


if __name__ == "__main__":
    unittest.main()
